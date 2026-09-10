"""生成《守望平台原创性声明》提交稿 docx。

排版规格复刻自 D:/揭榜挂帅/提交资料/总结报告.pdf 及其源 docx（总结报告.docx）:
  - A4; 页边距 上5.5cm/下3.5cm/左右3cm; 页眉页脚距 1.27cm
  - 正文: 仿宋(中文)/Times New Roman(西文) 12pt, 1.5倍行距, 首行缩进21pt(=420tw), 两端对齐
  - 条标题: 楷体 15pt; 内容页大标题: 黑体 22pt 居中
  - 封面: 黑体逐字竖排(52pt, 每字一段居中) + 副标题(黑体16pt, 首行缩进4字)
  - 表格: Table Grid 网格, 表头行楷体10.5pt, 数据单元格仿宋10.5pt
  - 页脚: 居中页码(PAGE 域, Cambria 11pt)

用法: python docs/competition-materials/build_declaration_docx.py
输入: 同目录 03-原创性声明.md (唯一内容源, 改完 md 重跑本脚本即可)
输出: 同目录 守望平台原创性声明.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent
MD_PATH = BASE / "03-原创性声明.md"
OUT_PATH = BASE / "守望平台原创性声明.docx"

LATIN, FANGSONG, HEITI, KAITI = "Times New Roman", "仿宋", "黑体", "楷体"
COVER_TITLE = "原创性声明"
COVER_SUBTITLE = "守望——基于多智能体编排的安全运营平台"


# ---------- 基础工具 ----------

def set_run(run, text, east=FANGSONG, size=12, bold=False):
    run.text = text
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    return run


def fmt_para(p, align=None, before=None, after=None, line=None, first=None, keep=False):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if first is not None:
        pf.first_line_indent = Pt(first)
    if keep:
        pf.keep_with_next = True
    return p


def split_bold(s):
    """把 **加粗** 拆成 [(text, bold), ...]"""
    parts = re.split(r"\*\*(.+?)\*\*", s)
    return [(t, i % 2 == 1) for i, t in enumerate(parts) if t]


def fill_runs(p, text, east=FANGSONG, size=12):
    for seg, bold in split_bold(text):
        set_run(p.add_run(), seg, east=east, size=size, bold=bold)


def add_para(doc, text, east=FANGSONG, size=12, **fmt):
    p = doc.add_paragraph()
    fill_runs(p, text, east=east, size=size)
    return fmt_para(p, **fmt)


# ---------- md 解析 ----------

def parse_md(path):
    blocks = []
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("# "):
            blocks.append(("title", s[2:].strip()))
        elif s.startswith("## "):
            blocks.append(("h2", s[3:].strip()))
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        elif re.match(r"^\d+\.\s", s):
            blocks.append(("li", s))
        else:
            blocks.append(("p", s))
        i += 1
    return blocks


# ---------- 组件 ----------

def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "8505")  # 15cm
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    grid_widths = ["2126", "6379"]  # 3.75cm + 11.25cm
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), grid_widths):
        gc.set(qn("w:w"), w)
    for ri, row in enumerate(rows):
        for ci, txt in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Cm(3.75 if ci == 0 else 11.25)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            fill_runs(p, txt, east=KAITI if ri == 0 else FANGSONG, size=10.5)
            fmt_para(p, before=2, after=2, line=1.0)
    return t


def build_footer(sec):
    xml = (
        f'<w:p {nsdecls("w")}>'
        '<w:pPr><w:jc w:val="center"/></w:pPr>'
        + "".join(
            f'<w:r><w:rPr><w:rFonts w:ascii="Cambria" w:hAnsi="Cambria"/><w:sz w:val="22"/></w:rPr>{body}</w:r>'
            for body in [
                '<w:fldChar w:fldCharType="begin"/>',
                '<w:instrText xml:space="preserve"> PAGE </w:instrText>',
                '<w:fldChar w:fldCharType="separate"/>',
                "<w:t>1</w:t>",
                '<w:fldChar w:fldCharType="end"/>',
            ]
        )
        + "</w:p>"
    )
    new_p = parse_xml(xml)
    old_p = sec.footer._element.find(qn("w:p"))
    sec.footer._element.replace(old_p, new_p)


# ---------- 主流程 ----------

def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FANGSONG)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(5.5), Cm(3.5)
    sec.left_margin = sec.right_margin = Cm(3)
    sec.header_distance = sec.footer_distance = Pt(36)

    # 封面: 逐字竖排标题 + 副标题 + 分页
    for ch in COVER_TITLE:
        p = doc.add_paragraph()
        set_run(p.add_run(), ch, east=HEITI, size=52)
        fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=20, after=20, line=1.15, keep=True)
    sub = doc.add_paragraph()
    set_run(sub.add_run(), COVER_SUBTITLE, east=HEITI, size=16)
    fmt_para(sub, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10, line=1.15, first=64.25)
    sub.add_run().add_break(WD_BREAK.PAGE)

    # 正文: 按 md 块渲染
    n_tables = 0
    for kind, val in parse_md(MD_PATH):
        if kind == "title":
            add_para(doc, val, east=HEITI, size=22,
                     align=WD_ALIGN_PARAGRAPH.CENTER, before=30, after=30, line=1.5, keep=True)
        elif kind == "h2":
            add_para(doc, val, east=KAITI, size=15, before=12, after=6, line=1.5, keep=True)
        elif kind == "li":
            add_para(doc, val, first=0, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "p":
            before = 12 if val.startswith("**附则**") else None
            add_para(doc, val, first=21, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=before)
        elif kind == "table":
            add_table(doc, val)
            n_tables += 1

    build_footer(sec)
    doc.core_properties.title = "守望平台·智能体原创性声明"
    doc.core_properties.author = "守望项目组"

    doc.save(OUT_PATH)
    print(f"OK -> {OUT_PATH}")
    print(f"blocks={len(parse_md(MD_PATH))} tables={n_tables}")


if __name__ == "__main__":
    main()
